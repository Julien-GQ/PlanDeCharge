from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')
ws = wb['CHARGE_HEBDO']

print("=== Vérification formatage CHARGE_HEBDO ===\n")

# Vérifier la ligne des semaines (table_header_row = 2)
print("Ligne des semaines (table_header_row):")
cell = ws['B2']
print(f"  Hauteur: {ws.row_dimensions[2].height}")
print(f"  Police couleur: {cell.font.color.rgb if cell.font and cell.font.color else 'None'}")
print(f"  Police gras: {cell.font.bold if cell.font else 'None'}")
print(f"  Police taille: {cell.font.size if cell.font else 'None'}")
print(f"  Fond couleur: {cell.fill.fgColor.rgb if cell.fill else 'None'}")
print(f"  Bordure: {cell.border.left.style if cell.border and cell.border.left else 'None'}")

# Vérifier les alignements
print(f"\nAlignement vertical (B2): {cell.alignment.vertical if cell.alignment else 'None'}")
print(f"Alignement horizontal (B2): {cell.alignment.horizontal if cell.alignment else 'None'}")

# Vérifier une cellule de données
print(f"\nCellule de données (B3):")
cell3 = ws['B3']
print(f"  Alignement vertical: {cell3.alignment.vertical if cell3.alignment else 'None'}")
print(f"  Bordure: {cell3.border.left.style if cell3.border and cell3.border.left else 'None'}")

# Vérifier CMD_STOCK aussi
ws_stock = wb['CMD_STOCK']
cell_stock = ws_stock['B2']
print(f"\n=== CMD_STOCK ===")
print(f"Ligne semaines - Hauteur: {ws_stock.row_dimensions[2].height}")
print(f"Ligne semaines - Fond: {cell_stock.fill.fgColor.rgb if cell_stock.fill else 'None'}")
