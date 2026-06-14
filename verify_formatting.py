from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')
ws = wb['CHARGE_HEBDO']

print("=== Vérification CHARGE_HEBDO ===")
print(f"Largeur colonne A: {ws.column_dimensions['A'].width}")
print(f"Largeur colonne B: {ws.column_dimensions['B'].width}")

print(f"\nCell A1: {ws['A1'].value}")
print(f"Cell A1 border: {ws['A1'].border.left.style if ws['A1'].border else 'None'}")

# Vérifier hauteur ligne et polices
print("\n=== Hauteur des lignes et polices ===")
for row_idx in [2, 3, 4, 5, 6]:
    height = ws.row_dimensions[row_idx].height
    cell_val = ws.cell(row_idx, 1).value
    font_size = ws.cell(row_idx, 1).font.size if ws.cell(row_idx, 1).font else 'None'
    print(f"Row {row_idx}: '{cell_val}' | Height: {height} | Font size: {font_size}")

# Vérifier les couleurs des cellules
print(f"\nCell B3 fill color (past week): {ws['B3'].fill.fgColor.rgb if ws['B3'].fill else 'None'}")
print(f"Cell B4 fill color: {ws['B4'].fill.fgColor.rgb if ws['B4'].fill else 'None'}")

# Vérifier bordures
print(f"\nCell B2 border left style: {ws['B2'].border.left.style if ws['B2'].border and ws['B2'].border.left else 'None'}")
print(f"Cell B2 border color: {ws['B2'].border.left.color.rgb if ws['B2'].border and ws['B2'].border.left and ws['B2'].border.left.color else 'None'}")
