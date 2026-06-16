from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')

print("=== Feuilles créées ===")
for sheet_name in wb.sheetnames:
    print(f"- {sheet_name}")

print("\n=== Détails feuille Retard ===")
if "Retard" in wb.sheetnames:
    ws = wb["Retard"]
    print(f"Nombre de lignes: {ws.max_row}")
    print(f"Nombre de colonnes: {ws.max_column}")
    
    print("\nRésumé (premières 5 lignes):")
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(3, ws.max_column + 1)):
            cell = ws.cell(row_idx, col_idx)
            row_data.append(str(cell.value))
        print(f"  Row {row_idx}: {' | '.join(row_data)}")
else:
    print("Aucune feuille 'Retard' trouvée")

print("\n=== Détails feuille Charge_Global ===")
if "Charge_Global" in wb.sheetnames:
    ws = wb["Charge_Global"]
    print(f"Nombre de lignes: {ws.max_row}")
    print(f"Nombre de colonnes: {ws.max_column}")
else:
    print("Aucune feuille 'Charge_Global' trouvée")
