from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')
ws = wb['CHARGE_HEBDO']

print("=== Vérification polices des données ===")
print("\nRows 3-6 (Type + données):")
for row_idx in [3, 4, 5, 6]:
    row_label = ws.cell(row_idx, 1).value
    print(f"\nRow {row_idx}: {row_label}")
    for col_idx in range(1, 5):
        cell = ws.cell(row_idx, col_idx)
        font_size = cell.font.size if cell.font else 'None'
        font_bold = cell.font.bold if cell.font else 'None'
        fill = cell.fill.fgColor.rgb if cell.fill else 'None'
        print(f"  Col {col_idx}: Font={font_size}pt bold={font_bold} | Fill={fill}")

# Vérifier la semaine courante
from datetime import date
iso = date.today().isocalendar()
print(f"\nSemaine actuelle: S{iso.week:02d}")
