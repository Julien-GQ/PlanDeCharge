from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')

print("=== Feuilles et tables ===")
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    tables = list(sheet.tables.values())
    print(f"\n{sheet_name}:")
    if tables:
        for table in tables:
            print(f"  - Table: {table.displayName}")
    else:
        print(f"  (pas de tables)")

print("\n✅ Fichier prêt à être testé")
