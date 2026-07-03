"""Pipeline configurable pour reorganiser les commandes en cours."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

try:
    from .ca_sheet import write_ca_sheet
    from .charge_sheet import write_weekly_summary_sheet
    from .navigation import apply_workbook_navigation
except ImportError:
    from importlib import import_module

    write_ca_sheet = import_module("ca_sheet").write_ca_sheet
    write_weekly_summary_sheet = import_module("charge_sheet").write_weekly_summary_sheet
    apply_workbook_navigation = import_module("navigation").apply_workbook_navigation


@dataclass
class TransformResult:
    """Resultat du traitement V1."""

    rows_total: int
    rows_kept: int
    output_path: Path
    output_sheet: str


@dataclass
class Profile:
    """Definition d'un profil de transformation JSON."""

    name: str
    expected_columns: list[str]
    computed_columns: dict[str, dict[str, Any]]
    filters: list[dict[str, Any]]
    output_sheet: str
    column_order: list[str]
    drop_columns: set[str]
    column_format: dict[str, dict[str, Any]]
    create_excel_table: bool
    sort_by_week: bool
    create_weekly_summary: bool
    summary_sheet_name: str


def _first_word(value: str) -> str:
    value = value.strip()
    if not value:
        return ""
    return value.split()[0]


def _parse_french_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_profile(path: Path) -> Profile:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)

    output = payload.get("output", {})
    return Profile(
        name=str(payload.get("profile", path.stem)),
        expected_columns=list(payload.get("expected_columns", [])),
        computed_columns=dict(payload.get("computed_columns", {})),
        filters=list(payload.get("filters", [])),
        output_sheet=str(output.get("sheet_name", "Resultat")),
        column_order=list(output.get("column_order", [])),
        drop_columns=set(output.get("drop_columns", [])),
        column_format=dict(output.get("column_format", {})),
        create_excel_table=bool(output.get("create_excel_table", True)),
        sort_by_week=bool(output.get("sort_by_week", True)),
        create_weekly_summary=bool(output.get("create_weekly_summary", True)),
        summary_sheet_name=str(output.get("summary_sheet_name", "Charge")),
    )


def _apply_filter_overrides(
    filters: list[dict[str, Any]],
    overrides: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    if not overrides:
        return list(filters)

    updated_filters = [dict(rule) for rule in filters]
    for column_name, value in overrides.items():
        replaced = False
        for rule in updated_filters:
            if str(rule.get("column", "")).strip().upper() != column_name.upper():
                continue
            if str(rule.get("operator", "")).strip().lower() != "equals":
                continue
            rule["value"] = value
            replaced = True
            break

        if not replaced:
            updated_filters.append(
                {"column": column_name, "operator": "equals", "value": value}
            )

    return updated_filters


def _extract_equals_filter_value(filters: list[dict[str, Any]], column_name: str) -> str | None:
    for rule in filters:
        if str(rule.get("column", "")).strip().upper() != column_name.upper():
            continue
        if str(rule.get("operator", "")).strip().lower() != "equals":
            continue
        value = _normalize_text(rule.get("value"))
        if value:
            return value
    return None


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _validate_structure(headers: list[str], expected_columns: list[str]) -> None:
    missing = [col for col in expected_columns if col not in headers]
    if missing:
        listed = ", ".join(missing)
        raise ValueError(f"Structure invalide: colonnes manquantes: {listed}")


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        with path.open("r", encoding="cp1252", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            headers = [h.strip() for h in (reader.fieldnames or [])]
            rows = []
            for raw_row in reader:
                row: dict[str, Any] = {}
                for key, value in raw_row.items():
                    if key is not None:
                        row[key.strip()] = value
                rows.append(row)
            return headers, rows

    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Format non supporte: {path.suffix}")

    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=(suffix == ".xlsm"))
    sheet = workbook.active
    for candidate in workbook.worksheets:
        if candidate.title.upper().startswith("DELAI_FAB_M"):
            sheet = candidate
            break
    values = sheet.iter_rows(values_only=True)
    first_row = next(values, None)
    if first_row is None:
        return [], []
    headers = [str(cell).strip() if cell is not None else "" for cell in first_row]
    rows = []
    for line in values:
        row = {headers[idx]: line[idx] for idx in range(len(headers)) if headers[idx]}
        rows.append(row)
    return headers, rows


def list_available_ateliers(input_path: Path) -> list[str]:
    headers, rows = _read_rows(input_path)
    if "MOR_ATELIER" not in headers:
        return []

    ateliers = {
        _normalize_text(row.get("MOR_ATELIER"))
        for row in rows
        if _normalize_text(row.get("MOR_ATELIER"))
    }
    return sorted(ateliers)


def _match_filter(value: Any, rule: dict[str, Any]) -> bool:
    text = _normalize_text(value)
    op = str(rule.get("operator", "")).lower()
    target = _normalize_text(rule.get("value"))

    if op == "equals":
        return text.upper() == target.upper()
    if op == "not_equals":
        if rule.get("allow_empty") and text == "":
            return True
        return text.upper() != target.upper()
    if op == "not_empty":
        return text != ""

    raise ValueError(f"Operateur de filtre non supporte: {op}")


def _is_row_kept(row: dict[str, Any], rules: list[dict[str, Any]]) -> bool:
    for rule in rules:
        col = str(rule.get("column", ""))
        if col == "":
            continue
        if not _match_filter(row.get(col), rule):
            return False
    return True


def _compute_value(row: dict[str, Any], rule: dict[str, Any]) -> Any:
    source = str(rule.get("source", ""))
    transform = str(rule.get("transform", ""))
    raw = row.get(source)

    if transform == "iso_week":
        dt = _parse_french_date(raw)
        return dt.isocalendar().week if dt else ""
    if transform == "first_word":
        return _first_word(_normalize_text(raw))

    raise ValueError(f"Transformation non supportee: {transform}")


def _transform_row(row: dict[str, Any], computed_columns: dict[str, dict[str, Any]]) -> dict[str, Any]:
    transformed = dict(row)
    for new_col, rule in computed_columns.items():
        transformed[new_col] = _compute_value(row, rule)
    return transformed


def _ordered_columns(
    headers: list[str],
    transformed_rows: list[dict[str, Any]],
    preferred_order: list[str],
    drop_columns: set[str],
) -> list[str]:
    seen = set()
    ordered = []

    for col in preferred_order:
        if any(col in row for row in transformed_rows):
            ordered.append(col)
            seen.add(col)

    for col in headers:
        if col and col not in seen and col not in drop_columns:
            ordered.append(col)
            seen.add(col)

    for row in transformed_rows:
        for col in row.keys():
            if col not in seen and col not in drop_columns:
                ordered.append(col)
                seen.add(col)

    return ordered


def _format_cell_value(column: str, value: Any, column_format: dict[str, dict[str, Any]]) -> Any:
    fmt = column_format.get(column, {})
    if fmt.get("date_without_time"):
        dt = _parse_french_date(value)
        if dt:
            return dt.strftime("%d/%m/%Y")
    return value


def _parse_quantity(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "")
    if text == "":
        return 0.0
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _is_ca_source_row(row: dict[str, Any], atelier_filter: str | None) -> bool:
    if atelier_filter:
        atelier = _normalize_text(row.get("MOR_ATELIER")).upper()
        if atelier != _normalize_text(atelier_filter).upper():
            return False

    client = _normalize_text(row.get("NOMREDCLI_CDE")).upper()
    if client in {"MORTELECQUE", "WILLMARK", "F-WILLMARK"}:
        return False

    fam_vente = _normalize_text(row.get("FAM_VENTE")).upper()
    if fam_vente in {"VTENEGG", "VTENEGLI"}:
        return False

    return True


def _week_key(row: dict[str, Any]) -> tuple[int, int] | None:
    dt_liv = _parse_french_date(row.get("DT_LIV_CONFIRMEE"))
    if dt_liv is None:
        return None
    iso = dt_liv.isocalendar()
    return (iso.year, iso.week)


def _week_label(week: tuple[int, int]) -> str:
    _, iso_week = week
    return f"S{iso_week:02d}"

def _write_retard_sheet(
    workbook: Any,
    sheet_name: str,
    transformed_rows: list[dict[str, Any]],
    columns: list[str],
    column_format: dict[str, dict[str, Any]],
    navigation_targets: list[str],
) -> None:
    """Cree la feuille Retard avec synthese hebdo et details OF."""
    current_iso = date.today().isocalendar()
    current_week_key = (current_iso.year, current_iso.week)

    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    sheet = workbook.create_sheet(sheet_name)
    _write_navigation_sidebar(sheet, navigation_targets)

    retard_rows: list[dict[str, Any]] = []
    for row in transformed_rows:
        week_key = _week_key(row)
        client_name = _normalize_text(row.get("NOMREDCLI_CDE")).upper()
        if week_key is not None and week_key < current_week_key and "WILLMARK" not in client_name:
            retard_rows.append(row)

    black_border = Border(
        left=Side(style="thin", color="FF000000"),
        right=Side(style="thin", color="FF000000"),
        top=Side(style="thin", color="FF000000"),
        bottom=Side(style="thin", color="FF000000"),
    )
    header_fill = PatternFill(fill_type="solid", fgColor="FF404040")

    # Bloc gauche: S / M / P / A par semaine + ligne TT
    for col_idx, label in enumerate(["S", "M", "P", "A"], start=3):
        cell = sheet.cell(1, col_idx, label)
        cell.font = Font(bold=True, size=11, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = black_border
        sheet.column_dimensions[get_column_letter(col_idx)].width = 8

    sheet.row_dimensions[1].height = 20

    week_matrix: dict[tuple[int, int], dict[str, float]] = defaultdict(
        lambda: {"M": 0.0, "P": 0.0, "A": 0.0}
    )
    for row in retard_rows:
        week_key = _week_key(row)
        if week_key is None:
            continue
        qty = _parse_quantity(row.get("RESTE_A_LIV_UV"))
        type_name = _normalize_text(row.get("TYPE")).upper()
        if type_name.startswith("MANCHE"):
            week_matrix[week_key]["M"] += qty
        elif type_name.startswith("POCHE"):
            week_matrix[week_key]["P"] += qty
        else:
            week_matrix[week_key]["A"] += qty

    ordered_weeks = sorted(week_matrix.keys())
    row_idx = 2
    total_m = 0.0
    total_p = 0.0
    total_a = 0.0

    for week in ordered_weeks:
        m_value = week_matrix[week]["M"]
        p_value = week_matrix[week]["P"]
        a_value = week_matrix[week]["A"]
        total_m += m_value
        total_p += p_value
        total_a += a_value

        values = [
            _week_label(week),
            int(round(m_value)) if m_value > 0 else "",
            int(round(p_value)) if p_value > 0 else "",
            int(round(a_value)) if a_value > 0 else "",
        ]
        for col_idx, value in enumerate(values, start=3):
            cell = sheet.cell(row_idx, col_idx, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = black_border
        row_idx += 1

    tt_row = row_idx
    tt_fill = PatternFill(fill_type="solid", fgColor="FFC0C0C0")
    tt_values = ["TT", int(round(total_m)), int(round(total_p)), int(round(total_a))]
    for col_idx, value in enumerate(tt_values, start=3):
        cell = sheet.cell(tt_row, col_idx, value)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = black_border
        cell.fill = tt_fill

    if tt_row >= 2:
        _apply_banded_rows(sheet, 2, tt_row - 1, 3, 6)
    _apply_table_gridlines(sheet, 1, tt_row, 3, 6)

    # Colonne de separation entre le bloc gauche et les details
    sheet.column_dimensions[get_column_letter(7)].width = 3

    # Bloc droite: details OF retard en conservant la mise en page de Global
    details_start_col = 8
    for offset, col_name in enumerate(columns):
        target_col = details_start_col + offset
        sheet.cell(1, target_col, col_name)

    for data_row_idx, row in enumerate(retard_rows, start=2):
        for offset, col_name in enumerate(columns):
            target_col = details_start_col + offset
            sheet.cell(
                data_row_idx,
                target_col,
                _format_cell_value(col_name, row.get(col_name, ""), column_format),
            )

    details_end_row = max(1, len(retard_rows) + 1)
    if columns:
        details_end_col = details_start_col + len(columns) - 1
        sheet.auto_filter.ref = (
            f"{get_column_letter(details_start_col)}1:{get_column_letter(details_end_col)}{details_end_row}"
        )

    for offset, col_name in enumerate(columns):
        target_col = details_start_col + offset
        fmt = column_format.get(col_name, {})

        width = fmt.get("width")
        if width is not None:
            sheet.column_dimensions[get_column_letter(target_col)].width = float(width)

        align_value = str(fmt.get("align", "")).lower()
        horizontal = None
        if align_value in {"center", "centre"}:
            horizontal = "center"
        elif align_value in {"left", "gauche"}:
            horizontal = "left"

        bold = bool(fmt.get("bold", False))
        font_size = fmt.get("font_size")
        font_size_value = float(font_size) if font_size is not None else None

        for row_idx in range(1, details_end_row + 1):
            cell = sheet.cell(row_idx, target_col)
            if horizontal:
                cell.alignment = Alignment(horizontal=horizontal, vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
            if bold or font_size_value is not None:
                cell.font = Font(bold=bold, size=font_size_value)
            cell.border = black_border

    if columns and details_end_row >= 2:
        _apply_banded_rows(sheet, 2, details_end_row, details_start_col, details_end_col)
    if columns:
        _apply_table_gridlines(sheet, 1, details_end_row, details_start_col, details_end_col)

    _apply_print_layout(sheet)


def _week_sort_key(row: dict[str, Any]) -> tuple[int, int, int, date]:
    week_raw = row.get("SEMAINE", "")
    dt_liv = _parse_french_date(row.get("DT_LIV_CONFIRMEE"))
    date_missing = 0 if dt_liv is not None else 1
    safe_date = dt_liv or date.max

    try:
        week = int(week_raw)
        return (0, week, date_missing, safe_date)
    except (TypeError, ValueError):
        return (1, 999, date_missing, safe_date)


def _apply_sheet_format(
    sheet: Any,
    columns: list[str],
    column_format: dict[str, dict[str, Any]],
    start_col: int = 1,
) -> None:
    for idx, col in enumerate(columns, start=start_col):
        fmt = column_format.get(col, {})
        width = fmt.get("width")
        if width:
            sheet.column_dimensions[get_column_letter(idx)].width = float(width)

        align_value = str(fmt.get("align", "")).lower()
        horizontal = None
        if align_value in {"center", "centre"}:
            horizontal = "center"
        elif align_value in {"left", "gauche"}:
            horizontal = "left"

        bold = bool(fmt.get("bold", False))
        font_size = fmt.get("font_size")
        font_size_value = float(font_size) if font_size is not None else None

        for row_idx in range(1, sheet.max_row + 1):
            cell = sheet.cell(row_idx, idx)
            if horizontal:
                cell.alignment = Alignment(horizontal=horizontal)
            if bold or font_size_value is not None:
                cell.font = Font(bold=bold, size=font_size_value)


def _apply_banded_rows(
    sheet: Any,
    start_row: int,
    end_row: int,
    start_col: int,
    end_col: int,
) -> None:
    """Applique un effet une ligne sur deux sur une zone tabulaire."""
    if end_row < start_row or end_col < start_col:
        return

    stripe_fill = PatternFill(fill_type="solid", fgColor="FFEAF2FB")
    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")

    for row_idx in range(start_row, end_row + 1):
        row_fill = stripe_fill if (row_idx - start_row) % 2 == 0 else white_fill
        for col_idx in range(start_col, end_col + 1):
            sheet.cell(row_idx, col_idx).fill = row_fill


def _apply_table_gridlines(
    sheet: Any,
    start_row: int,
    end_row: int,
    start_col: int,
    end_col: int,
) -> None:
    """Applique un quadrillage fin noir sur une zone tabulaire."""
    if end_row < start_row or end_col < start_col:
        return

    black_border = Border(
        left=Side(style="thin", color="FF000000"),
        right=Side(style="thin", color="FF000000"),
        top=Side(style="thin", color="FF000000"),
        bottom=Side(style="thin", color="FF000000"),
    )
    for row_idx in range(start_row, end_row + 1):
        for col_idx in range(start_col, end_col + 1):
            sheet.cell(row_idx, col_idx).border = black_border


def _set_uniform_row_height(sheet: Any, height: float = 18) -> None:
    """Uniformise la hauteur de toutes les lignes utilisees d'une feuille."""
    for row_idx in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row_idx].height = height


def _apply_print_layout(sheet: Any) -> None:
    # Mise en page orientee export PDF: 1 page en largeur, hauteur libre.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 9  # A4
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def _style_global_euro_column(sheet: Any, column_name: str = "€") -> None:
    """Applique le style attendu a la colonne € deja incluse dans le tableau."""
    col_idx = None
    for idx in range(1, sheet.max_column + 1):
        if str(sheet.cell(1, idx).value or "").strip() == column_name:
            col_idx = idx
            break

    if col_idx is None:
        return

    letter = get_column_letter(col_idx)
    pale_yellow = PatternFill(fill_type="solid", fgColor="FFFFF8DC")

    sheet.column_dimensions[letter].width = 10

    for row_idx in range(1, sheet.max_row + 1):
        cell = sheet.cell(row_idx, col_idx)
        if row_idx >= 2 and cell.value not in {None, ""}:
            cell.number_format = "# ##0"

        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(bold=True, size=11, color="FF000000")
        cell.fill = pale_yellow


def _write_navigation_sidebar(sheet: Any, sheet_names: list[str], width: float = 15) -> None:
    nav_fill = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
    nav_font = Font(bold=True, size=10, color="FF0563C1", underline="single")
    title_font = Font(bold=True, size=10, color="FF1F4E78")

    sheet.column_dimensions["A"].width = width
    sheet.cell(1, 1, "Navigation").font = title_font
    sheet.cell(1, 1).fill = nav_fill
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for name in sheet_names:
        link_cell = sheet.cell(row_idx, 1, name)
        link_cell.hyperlink = f"#{name}!A1"
        link_cell.font = nav_font
        link_cell.fill = nav_fill
        link_cell.alignment = Alignment(horizontal="left", vertical="center")
        row_idx += 1


def _write_like_global_sheet(
    workbook: Any,
    sheet_name: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    column_format: dict[str, dict[str, Any]],
    create_excel_table: bool,
    navigation_targets: list[str] | None = None,
    start_col: int = 3,
    apply_banding: bool = False,
) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    sheet = workbook.create_sheet(sheet_name)

    if navigation_targets:
        _write_navigation_sidebar(sheet, navigation_targets)

    for idx, col in enumerate(columns, start=start_col):
        sheet.cell(1, idx, col)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col in enumerate(columns, start=start_col):
            sheet.cell(row_idx, col_idx, _format_cell_value(col, row.get(col, ""), column_format))

    if create_excel_table and sheet.max_row >= 2 and sheet.max_column >= start_col:
        sheet.auto_filter.ref = (
            f"{get_column_letter(start_col)}1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )

    _apply_sheet_format(sheet, columns, column_format, start_col=start_col)
    if apply_banding and sheet.max_row >= 2 and sheet.max_column >= start_col:
        _apply_banded_rows(sheet, 2, sheet.max_row, start_col, sheet.max_column)
        _apply_table_gridlines(sheet, 1, sheet.max_row, start_col, sheet.max_column)
    _apply_print_layout(sheet)


def _write_current_week_charge_sheets(
    workbook: Any,
    columns: list[str],
    rows: list[dict[str, Any]],
    column_format: dict[str, dict[str, Any]],
    create_excel_table: bool,
    navigation_targets: list[str],
) -> None:
    current_iso = date.today().isocalendar()
    current_week_key = (current_iso.year, current_iso.week)

    current_week_rows = [row for row in rows if _week_key(row) == current_week_key]

    manches_rows: list[dict[str, Any]] = []
    poches_rows: list[dict[str, Any]] = []
    autres_rows: list[dict[str, Any]] = []

    for row in current_week_rows:
        type_name = _normalize_text(row.get("TYPE")).upper()
        if type_name.startswith("MANCHE"):
            manches_rows.append(row)
        elif type_name.startswith("POCHE"):
            poches_rows.append(row)
        else:
            autres_rows.append(row)

    _write_like_global_sheet(
        workbook,
        "S0_Manche",
        columns,
        manches_rows,
        column_format,
        create_excel_table,
        navigation_targets=navigation_targets,
        apply_banding=True,
    )
    _write_like_global_sheet(
        workbook,
        "S0_Poche",
        columns,
        poches_rows,
        column_format,
        create_excel_table,
        navigation_targets=navigation_targets,
        apply_banding=True,
    )
    _write_like_global_sheet(
        workbook,
        "S0_Autre",
        columns,
        autres_rows,
        column_format,
        create_excel_table,
        navigation_targets=navigation_targets,
        apply_banding=True,
    )


def _write_home_sheet(workbook: Any, sheet_names: list[str]) -> None:
    """Conservé pour compatibilite historique; Home n'est plus crée ici."""
    return


def _write_output(
    input_path: Path,
    output_path: Path,
    output_sheet: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    ca_atelier: str | None,
    column_format: dict[str, dict[str, Any]],
    create_excel_table: bool,
    create_weekly_summary: bool,
    summary_sheet_name: str,
) -> None:
    suffix = input_path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(input_path, keep_vba=(suffix == ".xlsm"))
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    # Nettoie les anciens noms de feuilles pour garder un classeur lisible.
    obsolete_sheets = [
        "Charge_Global",
        "CHARGE_HEBDO",
        "Charge_Semaine_Manches",
        "Charge_Semaine_Poches",
        "Charge_Semaine_Autres",
        "CA_details",
    ]
    for old_name in obsolete_sheets:
        if old_name in workbook.sheetnames:
            del workbook[old_name]

    home_targets = [
        output_sheet,
        summary_sheet_name,
        "CA",
        "CMD_STOCK",
        "Retard",
        "S0_Manche",
        "S0_Poche",
        "S0_Autre",
    ]

    global_columns = list(columns)
    global_rows = rows
    global_column_format = dict(column_format)
    if "€" not in global_columns:
        global_columns.append("€")
    global_rows = [
        dict(row, **{"€": int(round(_parse_quantity(row.get("CAPREV_LCDE")))) if abs(_parse_quantity(row.get("CAPREV_LCDE"))) > 1e-9 else ""})
        for row in rows
    ]
    global_column_format["€"] = {
        "width": 10,
        "align": "center",
        "bold": True,
        "font_size": 11,
    }

    _write_like_global_sheet(
        workbook,
        output_sheet,
        global_columns,
        global_rows,
        global_column_format,
        create_excel_table,
        navigation_targets=None,
    )
    _style_global_euro_column(workbook[output_sheet], "€")

    if create_weekly_summary:
        write_weekly_summary_sheet(workbook, summary_sheet_name, rows)

    rows_ca = [row for row in source_rows if _is_ca_source_row(row, ca_atelier)]
    write_ca_sheet(workbook, "CA", rows_ca)
    
    _write_retard_sheet(workbook, "Retard", rows, columns, column_format, home_targets)
    _write_current_week_charge_sheets(workbook, columns, rows, column_format, create_excel_table, home_targets)

    for sheet_name in home_targets:
        if sheet_name in workbook.sheetnames:
            _set_uniform_row_height(workbook[sheet_name], height=18)

    # Place la feuille principale generee en premier onglet pour l'affichage initial.
    if output_sheet in workbook.sheetnames:
        main_sheet = workbook[output_sheet]
        workbook._sheets.remove(main_sheet)
        workbook._sheets.insert(0, main_sheet)
        workbook.active = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    apply_workbook_navigation(output_path, target_sheet_names=home_targets)


def build_v1_sheet(
    input_path: Path,
    output_path: Path | None,
    profile_path: Path,
    output_sheet: str | None = None,
    filter_overrides: dict[str, Any] | None = None,
) -> TransformResult:
    """Genere une feuille de sortie depuis un profil JSON."""
    profile = _parse_profile(profile_path)
    headers, rows = _read_rows(input_path)
    _validate_structure(headers, profile.expected_columns)
    active_filters = _apply_filter_overrides(profile.filters, filter_overrides)
    ca_atelier = _extract_equals_filter_value(active_filters, "MOR_ATELIER")

    transformed_rows = [
        _transform_row(row, profile.computed_columns)
        for row in rows
        if _is_row_kept(row, active_filters)
    ]

    if profile.sort_by_week and any("SEMAINE" in row for row in transformed_rows):
        transformed_rows.sort(key=_week_sort_key)

    columns = _ordered_columns(
        headers,
        transformed_rows,
        preferred_order=profile.column_order,
        drop_columns=profile.drop_columns,
    )

    sheet_name = output_sheet or profile.output_sheet
    final_output = output_path
    if final_output is None:
        if input_path.suffix.lower() in {".xlsx", ".xlsm"}:
            final_output = input_path
        else:
            final_output = input_path.with_name(f"{input_path.stem}_resultat.xlsx")

    _write_output(
        input_path,
        final_output,
        sheet_name,
        columns,
        transformed_rows,
        rows,
        ca_atelier,
        profile.column_format,
        profile.create_excel_table,
        profile.create_weekly_summary,
        profile.summary_sheet_name,
    )

    return TransformResult(
        rows_total=len(rows),
        rows_kept=len(transformed_rows),
        output_path=final_output,
        output_sheet=sheet_name,
    )