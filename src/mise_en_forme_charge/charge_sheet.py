"""Generation de la feuille Charge (hebdomadaire)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_french_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_quantity(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(" ", "")
    if text == "":
        return 0.0
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _type_sort_key(type_name: str) -> tuple[int, str]:
    upper = type_name.upper()
    if upper.startswith("MANCHE"):
        return (0, upper)
    if upper.startswith("POCHE"):
        return (1, upper)
    return (2, upper)


def _week_key(row: dict[str, Any]) -> tuple[int, int] | None:
    dt_liv = _parse_french_date(row.get("DT_LIV_CONFIRMEE"))
    if dt_liv is None:
        return None
    iso = dt_liv.isocalendar()
    return (iso.year, iso.week)


def _week_label(week: tuple[int, int]) -> str:
    _, iso_week = week
    return f"S{iso_week:02d}"


def _friday_label(week: tuple[int, int]) -> str:
    year, iso_week = week
    friday = date.fromisocalendar(year, iso_week, 5)
    return friday.strftime("%d/%m")


def _build_weekly_matrix(
    rows: list[dict[str, Any]],
    weeks: list[tuple[int, int]],
) -> tuple[list[str], dict[str, dict[tuple[int, int], float]]]:
    matrix: dict[str, dict[tuple[int, int], float]] = defaultdict(lambda: defaultdict(float))
    types_set: set[str] = set()

    for row in rows:
        week = _week_key(row)
        if week is None:
            continue
        type_name = _normalize_text(row.get("TYPE")).upper()
        if type_name == "":
            type_name = "<SANS_TYPE>"
        qty = _parse_quantity(row.get("RESTE_A_LIV_UV"))
        matrix[type_name][week] += qty
        types_set.add(type_name)

    ordered_types = sorted(types_set, key=_type_sort_key)

    for type_name in ordered_types:
        for week in weeks:
            _ = matrix[type_name][week]

    return ordered_types, matrix


def _apply_print_layout(sheet: Any) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 9
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def _write_weekly_section(
    sheet: Any,
    start_row: int,
    rows: list[dict[str, Any]],
    weeks: list[tuple[int, int]],
    start_col: int = 3,
) -> int:
    current_iso = date.today().isocalendar()
    current_week_key = (current_iso.year, current_iso.week)
    yellow_pale = PatternFill(fill_type="solid", fgColor="FFFFF8DC")
    green_pale = PatternFill(fill_type="solid", fgColor="FFF0FFF0")
    black_border = Border(
        left=Side(style="thin", color="FF000000"),
        right=Side(style="thin", color="FF000000"),
        top=Side(style="thin", color="FF000000"),
        bottom=Side(style="thin", color="FF000000"),
    )

    friday_row = start_row
    table_header_row = start_row + 1
    data_start_row = start_row + 2

    sheet.cell(friday_row, start_col, "Date vendredi")
    sheet.cell(table_header_row, start_col, "TYPE")

    for idx, week in enumerate(weeks, start=start_col + 1):
        sheet.cell(friday_row, idx, _friday_label(week))
        sheet.cell(table_header_row, idx, _week_label(week))

    ordered_types, matrix = _build_weekly_matrix(rows, weeks)

    current_row = data_start_row
    for type_name in ordered_types:
        sheet.cell(current_row, start_col, type_name)
        emphasized = type_name.upper().startswith("MANCHE") or type_name.upper().startswith("POCHE")

        if emphasized:
            sheet.cell(current_row, start_col).font = Font(bold=True, size=12)
            sheet.row_dimensions[current_row].height = 20

        for idx, week in enumerate(weeks, start=start_col + 1):
            value = matrix[type_name][week]
            if abs(value) < 1e-9:
                sheet.cell(current_row, idx, "")
            else:
                cell_value = int(round(value))
                sheet.cell(current_row, idx, cell_value)
                sheet.cell(current_row, idx).number_format = "0"

            if emphasized:
                sheet.cell(current_row, idx).font = Font(bold=True, size=12)
            else:
                sheet.cell(current_row, idx).font = Font(bold=True, size=10)
        current_row += 1

    end_row = max(current_row - 1, table_header_row)
    end_col = start_col + len(weeks)

    sheet.column_dimensions[get_column_letter(start_col)].width = 23
    for idx in range(start_col + 1, end_col + 1):
        letter = get_column_letter(idx)
        sheet.column_dimensions[letter].width = 8

    for row_idx in range(friday_row, end_row + 1):
        sheet.cell(row_idx, start_col).alignment = Alignment(horizontal="left", vertical="center")
        for col_idx in range(start_col + 1, end_col + 1):
            sheet.cell(row_idx, col_idx).alignment = Alignment(horizontal="center", vertical="center")

    for row_idx in range(friday_row, end_row + 1):
        for col_idx in range(start_col, end_col + 1):
            sheet.cell(row_idx, col_idx).border = black_border

    for idx, week in enumerate(weeks, start=start_col + 1):
        if week < current_week_key:
            fill = yellow_pale
        elif week == current_week_key:
            fill = green_pale
        else:
            fill = None

        if fill is None:
            continue
        for row_idx in range(data_start_row, end_row + 1):
            sheet.cell(row_idx, idx).fill = fill

    for col_idx in range(start_col, end_col + 1):
        sheet.cell(friday_row, col_idx).font = Font(bold=False, size=10)
        header_cell = sheet.cell(table_header_row, col_idx)
        header_cell.font = Font(bold=True, size=11, color="FFFFFFFF")
        header_cell.fill = PatternFill(fill_type="solid", fgColor="FF404040")

    sheet.row_dimensions[table_header_row].height = 20
    return end_row + 2


def write_weekly_summary_sheet(
    workbook: Any,
    sheet_name: str,
    transformed_rows: list[dict[str, Any]],
) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    sheet = workbook.create_sheet(sheet_name)
    weeks = sorted({wk for row in transformed_rows if (wk := _week_key(row)) is not None})

    is_willmark = lambda row: "WILLMARK" in _normalize_text(row.get("NOMREDCLI_CDE")).upper()
    rows_main = [row for row in transformed_rows if not is_willmark(row)]
    rows_willmark = [row for row in transformed_rows if is_willmark(row)]

    _write_weekly_section(
        sheet,
        start_row=1,
        rows=rows_main,
        weeks=weeks,
        start_col=3,
    )

    if rows_willmark:
        stock_sheet_name = "CMD_STOCK"
        if stock_sheet_name in workbook.sheetnames:
            del workbook[stock_sheet_name]

        stock_sheet = workbook.create_sheet(stock_sheet_name)
        _write_weekly_section(
            stock_sheet,
            start_row=1,
            rows=rows_willmark,
            weeks=weeks,
            start_col=3,
        )

        _apply_print_layout(stock_sheet)

    _apply_print_layout(sheet)
