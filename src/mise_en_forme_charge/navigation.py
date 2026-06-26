"""Post-traitement de navigation pour les classeurs générés."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill


def _safe_sheet_order(sheet_names: list[str]) -> list[str]:
    preferred = ["Global", "Charge", "CA", "Retard", "S0_Manche", "S0_Poche", "S0_Autre"]
    ordered = [name for name in preferred if name in sheet_names]
    for name in sheet_names:
        if name not in ordered:
            ordered.append(name)
    return ordered


def _style_navigation_column(sheet: Any, sheet_names: list[str]) -> None:
    nav_fill = PatternFill(fill_type="solid", fgColor="FF0A2A66")
    nav_font = Font(bold=True, size=14, color="FFFFFFFF", underline="single")
    active_nav_font = Font(bold=True, size=14, color="FFFFE066", underline="single")
    nav_text_font = Font(bold=True, italic=True, size=16, color="FF7FB3E8")
    no_border = Border()

    # Nettoie une fusion precedente eventuelle avant de redefinir la zone.
    for merged_range in list(sheet.merged_cells.ranges):
        if str(merged_range) == "A1:A100":
            sheet.unmerge_cells("A1:A100")

    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 1

    nav_height = max(100, sheet.max_row + 5)
    for row_idx in range(1, nav_height + 1):
        nav_cell = sheet.cell(row_idx, 1)
        sep_cell = sheet.cell(row_idx, 2)

        nav_cell.fill = nav_fill
        nav_cell.border = no_border
        nav_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

        sep_cell.fill = nav_fill
        sep_cell.border = no_border
        sep_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Titre de menu
    title_cell = sheet.cell(1, 1, "Navigation")
    title_cell.font = nav_text_font
    title_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

    # Liens cliquables: un lien par cellule (Excel ne supporte pas plusieurs liens dans une cellule fusionnée).
    for row_idx, name in enumerate(sheet_names, start=3):
        if row_idx > nav_height:
            break
        link_cell = sheet.cell(row_idx, 1, f"• {name}")
        link_cell.hyperlink = f"#{name}!A1"
        link_cell.font = active_nav_font if name == sheet.title else nav_font
        link_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)


def apply_workbook_navigation(
    workbook_path: Path,
    target_sheet_names: list[str] | None = None,
) -> None:
    """Ajoute ou met a jour la navigation ancree sur les feuilles cible du classeur."""
    suffix = workbook_path.suffix.lower()
    workbook = load_workbook(workbook_path, keep_vba=(suffix == ".xlsm"))

    if "Home" in workbook.sheetnames:
        del workbook["Home"]

    if target_sheet_names:
        targets = [
            name for name in _safe_sheet_order(target_sheet_names)
            if name in workbook.sheetnames and name != "Home"
        ]
    else:
        targets = _safe_sheet_order([name for name in workbook.sheetnames if name != "Home"])

    for sheet_name in targets:
        _style_navigation_column(workbook[sheet_name], targets)

    workbook.save(workbook_path)
