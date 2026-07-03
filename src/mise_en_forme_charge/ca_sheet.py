"""Generation de la feuille CA (hebdomadaire et mensuelle)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _normalize_text(value)
    if text == "":
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = _normalize_text(value).replace(" ", "").replace(",", ".")
    if text == "":
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _first_word(value: str) -> str:
    value = value.strip()
    if not value:
        return ""
    return value.split()[0]


def _type_sort_key(type_name: str) -> tuple[int, str]:
    upper = type_name.upper()
    if upper.startswith("MANCHE"):
        return (0, upper)
    if upper.startswith("POCHE"):
        return (1, upper)
    return (2, upper)


def _week_key_from_row(row: dict[str, Any]) -> tuple[int, int] | None:
    dt_liv = _parse_date(row.get("DT_LIV_CONFIRMEE"))
    if dt_liv is None:
        return None
    iso = dt_liv.isocalendar()
    return (iso.year, iso.week)


def _month_key_from_row(row: dict[str, Any]) -> tuple[int, int] | None:
    dt_liv = _parse_date(row.get("DT_LIV_CONFIRMEE"))
    if dt_liv is None:
        return None
    return (dt_liv.year, dt_liv.month)


def _week_label(week: tuple[int, int]) -> str:
    return f"S{week[1]:02d}"


def _friday_label(week: tuple[int, int]) -> str:
    friday = date.fromisocalendar(week[0], week[1], 5)
    return friday.strftime("%d/%m")


def _month_label(month_key: tuple[int, int]) -> str:
    year, month = month_key
    return f"{month:02d}/{year}"


def _write_navigation_sidebar(sheet: Any, sheet_names: list[str], width: float = 15) -> None:
    nav_fill = PatternFill(fill_type="solid", fgColor="FFDDEBF7")
    nav_font = Font(bold=True, size=10, color="FF0563C1", underline="single")
    title_font = Font(bold=True, size=10, color="FF1F4E78")

    sheet.column_dimensions["A"].width = width
    sheet.cell(1, 1, "Navigation").font = title_font
    sheet.cell(1, 1).fill = nav_fill
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for name in sheet_names:
        link_cell = sheet.cell(row_idx, 1, name)
        link_cell.hyperlink = f"#{name}!A1"
        link_cell.font = nav_font
        link_cell.fill = nav_fill
        link_cell.alignment = Alignment(horizontal="left", vertical="center")
        row_idx += 1


def _aggregate_by_period(
    rows: list[dict[str, Any]],
    period_keys: list[tuple[int, int]],
    period_getter: Any,
) -> tuple[list[str], dict[str, dict[tuple[int, int], float]], dict[tuple[int, int], float]]:
    matrix: dict[str, dict[tuple[int, int], float]] = defaultdict(lambda: defaultdict(float))
    total_by_period: dict[tuple[int, int], float] = defaultdict(float)
    types: set[str] = set()

    for row in rows:
        period = period_getter(row)
        if period is None:
            continue
        raw_type = _normalize_text(row.get("TYPE"))
        if raw_type == "":
            raw_type = _first_word(_normalize_text(row.get("DESIGNATION")))
        type_name = raw_type.upper() or "<SANS_TYPE>"
        value = _parse_number(row.get("CAPREV_LCDE"))
        matrix[type_name][period] += value
        total_by_period[period] += value
        types.add(type_name)

    ordered_types = sorted(types, key=_type_sort_key)

    for type_name in ordered_types:
        for period in period_keys:
            _ = matrix[type_name][period]
    for period in period_keys:
        _ = total_by_period[period]

    return ordered_types, matrix, total_by_period


def _style_grid(
    sheet: Any,
    start_row: int,
    end_row: int,
    end_col: int,
    numeric_start_row: int,
    period_keys: list[tuple[int, int]],
    current_key: tuple[int, int],
    start_col: int = 1,
) -> None:
    black_border = Border(
        left=Side(style="thin", color="FF000000"),
        right=Side(style="thin", color="FF000000"),
        top=Side(style="thin", color="FF000000"),
        bottom=Side(style="thin", color="FF000000"),
    )
    yellow_pale = PatternFill(fill_type="solid", fgColor="FFFFF8DC")
    green_pale = PatternFill(fill_type="solid", fgColor="FFF0FFF0")

    for row_idx in range(start_row, end_row + 1):
        for col_idx in range(start_col, end_col + 1):
            cell = sheet.cell(row_idx, col_idx)
            cell.border = black_border
            if col_idx == start_col:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, key in enumerate(period_keys, start=start_col + 1):
        fill = None
        if key < current_key:
            fill = yellow_pale
        elif key == current_key:
            fill = green_pale
        if fill is None:
            continue
        for row_idx in range(numeric_start_row, end_row + 1):
            sheet.cell(row_idx, idx).fill = fill


def _write_weekly_section(sheet: Any, start_row: int, rows: list[dict[str, Any]], start_col: int = 1) -> int:
    weeks = sorted({wk for row in rows if (wk := _week_key_from_row(row)) is not None})
    end_col = start_col + len(weeks)

    friday_row = start_row
    header_row = start_row + 1
    total_row = start_row + 2
    data_start_row = start_row + 3

    sheet.cell(friday_row, start_col, "Date vendredi")
    sheet.cell(header_row, start_col, "TYPE")
    sheet.cell(total_row, start_col, "TOTAL SEMAINE")

    for idx, week in enumerate(weeks, start=start_col + 1):
        sheet.cell(friday_row, idx, _friday_label(week))
        sheet.cell(header_row, idx, _week_label(week))

    ordered_types, matrix, total_by_week = _aggregate_by_period(rows, weeks, _week_key_from_row)

    for idx, week in enumerate(weeks, start=start_col + 1):
        value = total_by_week[week]
        cell = sheet.cell(total_row, idx, int(round(value)) if abs(value) > 1e-9 else "")
        cell.number_format = "0"

    current_row = data_start_row
    for type_name in ordered_types:
        emphasized = type_name.startswith("MANCHE") or type_name.startswith("POCHE")
        type_cell = sheet.cell(current_row, start_col, type_name)
        type_cell.font = Font(bold=True, size=12 if emphasized else 10)

        for idx, week in enumerate(weeks, start=start_col + 1):
            value = matrix[type_name][week]
            val_cell = sheet.cell(current_row, idx, int(round(value)) if abs(value) > 1e-9 else "")
            val_cell.number_format = "0"
            val_cell.font = Font(bold=True, size=12 if emphasized else 10)
        current_row += 1

    end_row = max(current_row - 1, total_row)

    for col_idx in range(start_col, end_col + 1):
        sheet.cell(friday_row, col_idx).font = Font(bold=False, size=10)
        header_cell = sheet.cell(header_row, col_idx)
        header_cell.font = Font(bold=True, size=11, color="FFFFFFFF")
        header_cell.fill = PatternFill(fill_type="solid", fgColor="FF404040")

        total_cell = sheet.cell(total_row, col_idx)
        total_cell.font = Font(bold=True, size=11)
        total_cell.fill = PatternFill(fill_type="solid", fgColor="FFEDEDED")

    sheet.row_dimensions[header_row].height = 20
    sheet.column_dimensions[get_column_letter(start_col)].width = 23
    for idx in range(start_col + 1, end_col + 1):
        sheet.column_dimensions[get_column_letter(idx)].width = 8

    current_iso = date.today().isocalendar()
    current_week_key = (current_iso.year, current_iso.week)
    _style_grid(sheet, friday_row, end_row, end_col, data_start_row, weeks, current_week_key)

    return end_row + 2


def _write_monthly_section(sheet: Any, start_row: int, rows: list[dict[str, Any]], start_col: int = 1) -> int:
    months = sorted({mk for row in rows if (mk := _month_key_from_row(row)) is not None})
    end_col = start_col + len(months)

    header_row = start_row
    total_row = start_row + 1
    detail_sum_m_row = start_row + 2
    detail_sum_fw_row = start_row + 3
    detail_count_m_row = start_row + 4
    detail_count_fw_row = start_row + 5
    data_start_row = start_row + 6

    sheet.cell(header_row, start_col, "TYPE")
    sheet.cell(total_row, start_col, "TOTAL MOIS")
    sheet.cell(detail_sum_m_row, start_col, "CA : reste a livrer M")
    sheet.cell(detail_sum_fw_row, start_col, "CA : reste a livrer FW")
    sheet.cell(detail_count_m_row, start_col, "Nb cmd a livrer M")
    sheet.cell(detail_count_fw_row, start_col, "Nb cmd a livrer FW")

    for idx, month_key in enumerate(months, start=start_col + 1):
        sheet.cell(header_row, idx, _month_label(month_key))

    month_type_matrix: dict[str, dict[tuple[int, int], float]] = {
        "MANCHE": defaultdict(float),
        "POCHE": defaultdict(float),
        "AUTRE": defaultdict(float),
    }
    total_by_month: dict[tuple[int, int], float] = defaultdict(float)

    cmd_m_count_by_month: dict[tuple[int, int], int] = defaultdict(int)
    cmd_m_sum_by_month: dict[tuple[int, int], float] = defaultdict(float)
    cmd_m_seen_by_month: dict[tuple[int, int], set[str]] = defaultdict(set)

    cmd_fw_count_by_month: dict[tuple[int, int], int] = defaultdict(int)
    cmd_fw_sum_by_month: dict[tuple[int, int], float] = defaultdict(float)
    cmd_fw_seen_by_month: dict[tuple[int, int], set[str]] = defaultdict(set)

    for row in rows:
        month_key = _month_key_from_row(row)
        if month_key is None:
            continue
        caprev_value = _parse_number(row.get("CAPREV_LCDE"))
        total_by_month[month_key] += caprev_value

        raw_type = _normalize_text(row.get("TYPE"))
        if raw_type == "":
            raw_type = _first_word(_normalize_text(row.get("DESIGNATION")))
        type_name = raw_type.upper()
        if type_name.startswith("MANCHE"):
            bucket = "MANCHE"
        elif type_name.startswith("POCHE"):
            bucket = "POCHE"
        else:
            bucket = "AUTRE"
        month_type_matrix[bucket][month_key] += caprev_value

        commande = _normalize_text(row.get("COMMANDE"))
        if len(commande) == 5:
            if commande not in cmd_m_seen_by_month[month_key]:
                cmd_m_seen_by_month[month_key].add(commande)
                cmd_m_count_by_month[month_key] += 1
            cmd_m_sum_by_month[month_key] += caprev_value
        else:
            if commande and commande not in cmd_fw_seen_by_month[month_key]:
                cmd_fw_seen_by_month[month_key].add(commande)
                cmd_fw_count_by_month[month_key] += 1
            cmd_fw_sum_by_month[month_key] += caprev_value

    for idx, month_key in enumerate(months, start=start_col + 1):
        value = total_by_month[month_key]
        cell = sheet.cell(total_row, idx, int(round(value)) if abs(value) > 1e-9 else "")
        cell.number_format = "# ##0"

        count_m_cell = sheet.cell(detail_count_m_row, idx, cmd_m_count_by_month[month_key] or "")
        count_m_cell.number_format = "0"

        sum_m_value = cmd_m_sum_by_month[month_key]
        sum_m_cell = sheet.cell(detail_sum_m_row, idx, int(round(sum_m_value)) if abs(sum_m_value) > 1e-9 else "")
        sum_m_cell.number_format = "# ##0"

        count_fw_cell = sheet.cell(detail_count_fw_row, idx, cmd_fw_count_by_month[month_key] or "")
        count_fw_cell.number_format = "0"

        sum_fw_value = cmd_fw_sum_by_month[month_key]
        sum_fw_cell = sheet.cell(detail_sum_fw_row, idx, int(round(sum_fw_value)) if abs(sum_fw_value) > 1e-9 else "")
        sum_fw_cell.number_format = "# ##0"

    ordered_types = ["MANCHE", "POCHE", "AUTRE"]

    current_row = data_start_row
    for type_name in ordered_types:
        emphasized = type_name.startswith("MANCHE") or type_name.startswith("POCHE")
        type_cell = sheet.cell(current_row, start_col, type_name)
        type_cell.font = Font(bold=True, size=12 if emphasized else 10)

        for idx, month_key in enumerate(months, start=start_col + 1):
            value = month_type_matrix[type_name][month_key]
            val_cell = sheet.cell(current_row, idx, int(round(value)) if abs(value) > 1e-9 else "")
            val_cell.number_format = "# ##0"
            val_cell.font = Font(bold=True, size=12 if emphasized else 10)
        current_row += 1

    end_row = max(current_row - 1, detail_sum_fw_row)

    for col_idx in range(start_col, end_col + 1):
        header_cell = sheet.cell(header_row, col_idx)
        header_cell.font = Font(bold=True, size=11, color="FFFFFFFF")
        header_cell.fill = PatternFill(fill_type="solid", fgColor="FF404040")

        total_cell = sheet.cell(total_row, col_idx)
        total_cell.font = Font(bold=True, size=12)
        total_cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")

        detail_count_m_cell = sheet.cell(detail_count_m_row, col_idx)
        detail_count_m_cell.font = Font(bold=False, size=10)
        detail_count_m_cell.fill = PatternFill(fill_type="solid", fgColor="FFFFF8DC")

        detail_sum_m_cell = sheet.cell(detail_sum_m_row, col_idx)
        detail_sum_m_cell.font = Font(bold=True, size=10)
        detail_sum_m_cell.fill = PatternFill(fill_type="solid", fgColor="FFCCFFCC")

        detail_count_fw_cell = sheet.cell(detail_count_fw_row, col_idx)
        detail_count_fw_cell.font = Font(bold=False, size=10)
        detail_count_fw_cell.fill = PatternFill(fill_type="solid", fgColor="FFFFF8DC")

        detail_sum_fw_cell = sheet.cell(detail_sum_fw_row, col_idx)
        detail_sum_fw_cell.font = Font(bold=True, size=10)
        detail_sum_fw_cell.fill = PatternFill(fill_type="solid", fgColor="FFCCFFCC")

    sheet.row_dimensions[header_row].height = 20
    sheet.row_dimensions[total_row].height = 22
    sheet.row_dimensions[detail_sum_m_row].height = 20
    sheet.row_dimensions[detail_sum_fw_row].height = 20
    sheet.row_dimensions[detail_count_m_row].height = 20
    sheet.row_dimensions[detail_count_fw_row].height = 20

    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for row_idx in range(data_start_row, end_row + 1):
        sheet.row_dimensions[row_idx].height = 20
        for col_idx in range(start_col, end_col + 1):
            sheet.cell(row_idx, col_idx).fill = white_fill

    sheet.column_dimensions[get_column_letter(start_col)].width = 23
    for idx in range(start_col + 1, end_col + 1):
        sheet.column_dimensions[get_column_letter(idx)].width = 9

    current_today = date.today()
    current_month_key = (current_today.year, current_today.month)
    _style_grid(
        sheet,
        header_row,
        end_row,
        end_col,
        data_start_row,
        months,
        current_month_key,
        start_col=start_col,
    )

    return end_row


def _apply_print_layout(sheet: Any) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = 9
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def write_ca_sheet(workbook: Any, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    sheet = workbook.create_sheet(sheet_name)

    next_row = _write_monthly_section(sheet, 1, rows, start_col=3)
    _write_weekly_section(sheet, next_row + 2, rows, start_col=3)

    _apply_print_layout(sheet)
