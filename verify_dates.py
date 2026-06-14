from openpyxl import load_workbook

wb = load_workbook('base/DELAI_FAB_MJULIEN.xlsm')
ws_charge = wb['CHARGE_HEBDO']
ws_stock = wb['CMD_STOCK']

print("=== Vérification dates et polices ===\n")

for sheet_name, ws in [('CHARGE_HEBDO', ws_charge), ('CMD_STOCK', ws_stock)]:
    print(f"{sheet_name}:")
    friday_row = 1  # friday_row is start_row = 1
    
    # Vérifier les dates et polices
    for col_idx in range(2, 6):
        cell = ws.cell(friday_row, col_idx)
        date_value = cell.value
        font_size = cell.font.size if cell.font else 'None'
        print(f"  Col {col_idx}: '{date_value}' | Font: {font_size}pt")
    
    print()
